# -*- coding: utf-8 -*-
"""THE PRINCIPLE 方向性たたき v1 — PPTX generator (python-pptx)"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = "20242C"; PAPER = "FFFFFF"; SOFT = "5A6070"
BLUE = "1F4FA3"; BLUESOFT = "E8EEF8"; RED = "B0442C"; REDSOFT = "F7ECE8"
LINE = "D9D9D2"; SURFACE = "F2F2EE"; ICEBLUE = "7AA3E8"; GRAYD = "C9CDD6"; DKLINE = "3A3E48"
MINCHO = "Yu Mincho"; GOTHIC = "Yu Gothic"

W, H, MX = 13.333, 7.5, 0.7
CW = W - MX * 2

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

def rgb(hexs): return RGBColor.from_string(hexs)

def set_run_font(run, name, size, bold, color, spc=None):
    f = run.font
    f.name = name; f.size = Pt(size); f.bold = bold
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)
    if spc: rPr.set('spc', str(spc))

def tb(slide, x, y, w, h, content, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
       font=GOTHIC, anchor=MSO_ANCHOR.TOP, spacing=None, fill=None, border=None,
       spc=None, space_after=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if fill:
        box.fill.solid(); box.fill.fore_color.rgb = rgb(fill)
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = tf.margin_bottom = Inches(0.12)
    if border:
        box.line.color.rgb = rgb(border[0]); box.line.width = Pt(border[1])
    # content: str, or list of paragraphs; paragraph: str or list of (text, overrides)
    if isinstance(content, str):
        paras = content.split("\n")
    else:
        paras = content
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if spacing: p.line_spacing = Pt(spacing)
        if space_after: p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for txt, ov in runs:
            r = p.add_run(); r.text = txt
            set_run_font(r, ov.get("font", font), ov.get("size", size),
                         ov.get("bold", bold), ov.get("color", color), spc=ov.get("spc", spc))
    return box

def rect(slide, x, y, w, h, fill=None, border=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if border: sh.line.color.rgb = rgb(border[0]); sh.line.width = Pt(border[1])
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def hline(slide, x, y, w, color=LINE, pt=1):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(pt)
    return ln

def base(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else PAPER)
    return s

def header(s, no, title, sub=None):
    tb(s, MX, 0.40, CW, 0.3, no, size=10.5, bold=True, color=RED, font="Arial", spc=300)
    tb(s, MX, 0.70, CW, 0.75, title, size=26, bold=True, color=INK, font=MINCHO)
    if sub:
        tb(s, MX, 1.42, CW, 0.35, sub, size=12, color=SOFT)
    tb(s, MX, 7.10, CW, 0.28, "『THE PRINCIPLE』（仮）方向性ディスカッション — 翔泳社 押久保・大久保 / CONFIDENTIAL",
       size=8.5, color=SOFT)
    return s

def pill(slide, x, y, w, txt, bg, fg="FFFFFF"):
    sh = rect(slide, x, y, w, 0.32, fill=bg)
    tf = sh.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt
    set_run_font(r, "Arial", 10, True, fg, spc=150)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

def styled_table(slide, x, y, rows, col_w, row_h, header_fill=INK, hd_col=None,
                 fs_head=11.5, fs_body=10.5, special_fills=None):
    """rows: list of list of str. hd_col: index of column to bold w/ SURFACE fill.
    special_fills: dict {(r,c): (fill,color,bold)}"""
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_r))
    table = shp.table
    table.first_row = False; table.horz_banding = False
    tbl = table._tbl
    # remove theme style
    for el in tbl.findall(qn('a:tblPr')):
        for st in el.findall(qn('a:tableStyleId')): el.remove(st)
    for i, wd in enumerate(col_w): table.columns[i].width = Inches(wd)
    for i in range(n_r): table.rows[i].height = Inches(row_h)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill_c, col_c, bold_c = PAPER, INK, False
            if ri == 0: fill_c, col_c, bold_c = header_fill, "FFFFFF", True
            elif hd_col is not None and ci == hd_col: fill_c, bold_c = SURFACE, True
            if special_fills and (ri, ci) in special_fills:
                fill_c, col_c, bold_c = special_fills[(ri, ci)]
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill_c)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.line_spacing = Pt(14 if ri else 15)
            r = p.add_run(); r.text = val
            set_run_font(r, GOTHIC, fs_head if ri == 0 else fs_body, bold_c, col_c)
    return table

# ============ 1. Title ============
s = base(dark=True)
tb(s, MX, 1.25, CW, 0.35, "SHOEISHA  EDITORIAL DISCUSSION DRAFT — CONFIDENTIAL",
   size=11, bold=True, color=ICEBLUE, font="Arial", spc=400)
tb(s, MX, 1.9, CW, 2.3, "『THE PRINCIPLE』（仮）\n方向性ディスカッション",
   size=42, bold=True, color="FFFFFF", font=MINCHO, spacing=58)
tb(s, MX, 4.35, CW, 0.5, "骨子ドラフトを受けた「軸・見せ方・伝える順番」のたたき台",
   size=17, color=GRAYD)
hline(s, MX, 5.35, CW, color=DKLINE, pt=1)
tb(s, MX, 5.55, CW, 0.4, [[
    ("2026年8月28日　", {"color": "FFFFFF"}),
    ("翔泳社 押久保剛・大久保　｜　宛先：福田康隆さん MTG（1.5h）　｜　素材：骨子ドラフト（序章・第1〜4章・終章）＋7月メモ", {"color": "9BA1AE"}),
]], size=12)

# ============ 2. Goal ============
s = header(base(), "00 — GOAL", "本日のゴール：90分で3つ決める")
goals = [
    ("1", "読者と軸", "誰の・どんな場面の本か。「原理原則」をどの角度から売るか", "本日決めたい"),
    ("2", "構成の背骨", "章立ての基本方針（続編ファースト型か、フェーズ背骨型か）", "方向感まで"),
    ("3", "進め方", "執筆順序・レビューサイクル・タイトル検討の段取り", "本日決めたい"),
]
y = 2.0
for g in goals:
    rect(s, MX, y, CW, 1.06, fill=SURFACE, border=(LINE, 0.75))
    tb(s, MX + 0.22, y + 0.10, 0.8, 0.86, g[0], size=34, bold=True, color=BLUE, font=MINCHO, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, MX + 1.05, y + 0.14, 2.6, 0.78, g[1], size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, MX + 3.8, y + 0.14, 6.4, 0.78, g[2], size=12.5, anchor=MSO_ANCHOR.MIDDLE)
    is_dir = g[3] == "方向感まで"
    box = rect(s, MX + 10.35, y + 0.30, 1.45, 0.46, fill=BLUESOFT if is_dir else REDSOFT)
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = g[3]
    set_run_font(r, GOTHIC, 11, True, BLUE if is_dir else RED)
    y += 1.26
tb(s, MX, y + 0.1, CW, 0.4, "タイムボックス案：前提共有 10分 → 読者と軸 30分 → 構成 30分 → 見せ方・タイトル 10分 → 進め方 10分",
   size=12, color=SOFT)

# ============ 3. Summary ============
s = header(base(), "SUMMARY", "編集部の推し（先に結論）", "当日はこの4点を叩いてもらう前提のたたき台です")
recs = [
    ("読者", "THE MODEL読者の「7年後」", "当時の営業・マーケ実務層が、いま経営を任されている。続編の読者は新しい誰かではなく、同じ読者のいまの悩みに定める。"),
    ("軸", "原理原則を売り物に、フェーズを背骨に", "最も新しい発見は「フェーズの誤認が最大の失敗要因」（19社を横から見た帰結）。原理原則はフェーズという地図に載せた瞬間、現在地から引ける実用書になる。"),
    ("順番", "THE MODELアンサーは1章に圧縮", "過去の答え合わせで前半を使い切らず、「横の視点 → フェーズの発見」をできるだけ早く出す。"),
    ("タイトル", "『THE PRINCIPLE』を推す", "序章の結び「自分にとってのPRINCIPLEを作ってほしい」と完全整合。単数形の意味（ダリオとの差別化）も当日議論したい。"),
]
cw2 = (CW - 0.3) / 2
for i, it in enumerate(recs):
    x = MX + (i % 2) * (cw2 + 0.3)
    y = 2.0 + (i // 2) * 2.5
    rect(s, x, y, cw2, 2.3, fill=SURFACE, border=(LINE, 0.75))
    pill(s, x + 0.25, y + 0.25, 1.1, it[0], BLUE)
    tb(s, x + 0.25, y + 0.68, cw2 - 0.5, 0.45, it[1], size=16, bold=True, font=MINCHO)
    tb(s, x + 0.25, y + 1.18, cw2 - 0.5, 1.0, it[2], size=11.5, spacing=17)

# ============ 4. Draft map ============
s = header(base(), "01 — CURRENT DRAFT", "骨子ドラフトの現状マップ",
           "共有6ファイルを読み込み整理。「P」印の原理原則の断片は、そのまま商品価値になる粒が揃っている")
rows = [
    ["章", "内容"],
    ["序章", "なぜもう一度書くのか — 環境の変化（SaaS is Dead／AI）＋立場の変化（縦→横の視点）。原理原則＝判断に根拠を持つための蓄積。マシン／経営視点の定義"],
    ["第1章", "「THE MODEL」とはなんだったのか — 成立秘話（vs シーベル、2004年スライド）／よくある疑問への回答／7年の変化（AIエージェント、支援から完結へ）"],
    ["第2章", "横から見るJapan Cloudと「フェーズ」という発見 — 持続的成長の条件／成長4パターン／時間差の原則　※後半にGTM・報酬設計・マネジメントシステムまで同居＝過積載"],
    ["第3章", "フェーズ別の原理原則 — STARS 5フェーズ×（特徴／重点アクション／原理原則）。「フェーズの誤認が最大の失敗要因」"],
    ["第4章", "採用・リーダーシップ・キャリア — マルチプライヤー／やってはいけない5つ・やるべき4つ（新任リーダーの90日論）"],
    ["終章", "自分のPRINCIPLESを作る — レゴと粘土／各章ワークのとりまとめ／Journey is the reward（退職メール全文）"],
]
styled_table(s, MX, 2.05, rows, [1.3, CW - 1.3], 0.68, hd_col=0)

# ============ 5. Structural signs ============
s = header(base(), "01 — CURRENT DRAFT", "ドラフトから読み取れる構造上のサイン",
           "いま構成を決めるのが最も効率的なタイミング")
signs = [
    ("章構成は未確定", "本文中に第5章（マネジメントシステム）・第6章（立ち上げ）・第7章（新任リーダー）への参照があり、頭の中の構成は7章前後。共有ドラフトの4章立てとズレている。"),
    ("第2章は分割必至", "「フェーズという発見」「GTM原理原則」「マネジメントシステム」が1章に同居。どう割るかの判断が、そのまま「本の軸」の議論になる。"),
    ("体験設計の種がある", "各章末ワーク＋終章での「自分のPRINCIPLE」とりまとめ。読者が書き込む一冊通しの体験は、前作にない差別化資産。"),
]
cw3 = (CW - 0.6) / 3
for i, it in enumerate(signs):
    x = MX + i * (cw3 + 0.3)
    rect(s, x, 2.15, cw3, 3.7, fill=SURFACE, border=(LINE, 0.75))
    tb(s, x + 0.28, 2.4, 1, 0.6, str(i + 1), size=30, bold=True, color=RED, font=MINCHO)
    tb(s, x + 0.28, 3.05, cw3 - 0.56, 0.65, it[0], size=16, bold=True, font=MINCHO)
    tb(s, x + 0.28, 3.75, cw3 - 0.56, 1.9, it[1], size=11.5, spacing=17.5)

# ============ 6. Premise ============
s = header(base(), "02 — PREMISE", "前提の確認：勝ち筋と、最大の敵")
rect(s, MX, 2.0, cw2, 3.35, fill=BLUESOFT, border=(BLUE, 1))
tb(s, MX + 0.3, 2.3, cw2 - 0.6, 0.5, "勝ち筋は「誰が言うか」", size=18, bold=True, color=BLUE, font=MINCHO)
tb(s, MX + 0.3, 2.9, cw2 - 0.6, 2.3,
   "原理原則は誰でも語れる。しかし「累計19社の日本法人設立に関わり、12社の取締役として同じ時期に横から見た」人は日本にほぼいない。この観測点の希少性が本書の根拠。序章の「横の視点」がそれ。",
   size=12.5, spacing=20)
x2 = MX + cw2 + 0.3
rect(s, x2, 2.0, cw2, 3.35, fill=REDSOFT, border=(RED, 1))
tb(s, x2 + 0.3, 2.3, cw2 - 0.6, 0.5, "最大の敵は「原理原則」の抽象度", size=18, bold=True, color=RED, font=MINCHO)
tb(s, x2 + 0.3, 2.9, cw2 - 0.6, 2.3,
   "THE MODELが売れたのは、持ち帰れる「型」があったから。本書は「型ではなく原則だ」と言う本であり、型を否定する本は型ほど売れにくい。だから編集の仕事は、型の代わりの「持ち帰れる道具」（フェーズ地図・原則の型化・ワーク）を設計すること。",
   size=12.5, spacing=20)
tb(s, MX, 5.6, CW, 1.25, [[
    ("あえて指摘：", {"bold": True, "color": RED}),
    ("狙いは「フレームワーク紹介で終わらせない」なのに、現骨子はSTARS・Playing to Win・ダリオ等、他者フレームワークの引用が背骨になりつつある。終章の「レゴと粘土」こそ本書の思想 — この思想を序盤に宣言し、主役は福田さん自身の「P」、フレームワークは脇役に。", {}),
]], size=12, spacing=18, fill=SURFACE)

# ============ 7. Issue 1 readers ============
s = header(base(), "03 — ISSUE 1", "論点1：読者は誰か")
rows = [
    ["案", "読者像", "強み", "弱み"],
    ["a", "THE MODEL読者の「7年後」— 当時の実務層がいま部長・役員・社長に", "前作の読者基盤を引き継げる。悩みの変化（実行→設計）が本の主題と一致", "読者の高齢化とともに市場が狭まる"],
    ["b", "これから経営を担う人 — GM候補・次世代リーダー・経営者", "市場が広い。「体系的に経営を学ぶ場が日本にない」への回答になる", "経営書の棚は激戦区。続編の意味が弱まる"],
    ["c", "SaaS/IT業界の当事者", "具体性が最も活きる。講演・SNSの初速", "市場が最も狭い。「業界本」に見えると前作の業界外読者を失う"],
]
sp = {(1, 0): (SURFACE, BLUE, True), (2, 0): (SURFACE, BLUE, True), (3, 0): (SURFACE, BLUE, True)}
styled_table(s, MX, 2.0, rows, [0.7, 4.7, 3.7, 2.83], 0.82, special_fills=sp)
tb(s, MX, 5.6, CW, 1.15, [[
    ("推し：a を主読者、b を拡張読者に。", {"bold": True, "color": BLUE}),
    ("「THE MODELで実行を学んだあなたが、今度は設計する側に回る番だ」という一本のメッセージで a→b が繋がる。帯・序章の語りかけの主語をここに固定したい。", {}),
]], size=12.5, spacing=19, fill=BLUESOFT, border=(BLUE, 1))

# ============ 8. Issue 2 axes ============
s = header(base(), "04 — ISSUE 2", "論点2：本の軸 — 3つの候補",
           "福田さんの「どういう軸で書くのがいいか迷っている」への選択肢")
axes = [
    ("軸A", "続編軸 —「THE MODEL、その後」", "前作の誤解を解き、7年の変化（SaaS is Dead／AI）に答える。",
     "入口として最強。ただし軸にすると「答え合わせ本」で終わる。→ 1章に圧縮して入口に使う。", BLUE),
    ("軸B", "原理原則軸 —「福田版PRINCIPLES」", "判断に根拠を持つための蓄積を、エピソードとともに編む。終章が着地。",
     "本書のテーマそのもの。ただし並べるだけでは総花的。売り物ではあるが、背骨は別に必要。", BLUE),
    ("軸C", "フェーズ軸 —「現在地の地図」", "19社を横から見た発見＝フェーズの誤認が最大の失敗要因。フェーズごとに効く原則が変わる。",
     "最も新しく、最も福田さんにしか書けない主張。読者が現在地から本を引ける実用の背骨。", RED),
]
for i, it in enumerate(axes):
    x = MX + i * (cw3 + 0.3)
    rect(s, x, 2.1, cw3, 3.55, fill=SURFACE, border=(LINE, 0.75))
    pill(s, x + 0.25, 2.35, 0.95, it[0], it[4])
    tb(s, x + 0.25, 2.8, cw3 - 0.5, 0.8, it[1], size=14, bold=True, font=MINCHO, spacing=19)
    tb(s, x + 0.25, 3.62, cw3 - 0.5, 0.95, it[2], size=10.5, color=SOFT, spacing=15.5)
    tb(s, x + 0.25, 4.6, cw3 - 0.5, 1.0, it[3], size=10.5, spacing=15.5)
tb(s, MX, 5.9, CW, 1.0, [[
    ("推し：Cを背骨に、Bを売り物に、Aは第1章へ圧縮。", {"bold": True, "color": BLUE}),
    ("一文で「会社のフェーズを見誤るな。フェーズごとに効く原理原則を、19社の実話で渡す」。前作＝分業という空間の地図、本作＝フェーズという時間の地図。", {}),
]], size=12, spacing=18, fill=BLUESOFT, border=(BLUE, 1))

# ============ 9. Issue 3 structure ============
s = header(base(), "05 — ISSUE 3", "論点3：伝える順番 — 構成2案の比較")
rect(s, MX, 2.0, cw2, 3.2, fill=SURFACE, border=(LINE, 0.75))
pill(s, MX + 0.25, 2.2, 0.95, "案 1", INK)
tb(s, MX + 0.25, 2.6, cw2 - 0.5, 0.4, "続編ファースト型（現骨子の整流）", size=15.5, bold=True, font=MINCHO)
tb(s, MX + 0.25, 3.08, cw2 - 0.5, 1.0,
   "序章 → THE MODELアンサー → 横の視点と時間差 → GTM戦略の原則 → マネジメントシステム → フェーズ別原則 → 採用・リーダーシップ → 終章",
   size=11, color=SOFT, spacing=16.5)
tb(s, MX + 0.25, 4.12, cw2 - 0.5, 1.0, [
    [("○ 現ドラフトから距離が近く執筆負荷が小さい。前作読者の期待に素直", {"color": BLUE})],
    [("× テーマ別の並びで「どこを読めばいいか」が見えにくい。前半が過去の話で重い", {"color": RED})],
], size=11, spacing=16)
x2 = MX + cw2 + 0.3
rect(s, x2, 2.0, cw2, 3.2, fill=PAPER, border=(RED, 1.25))
pill(s, x2 + 0.25, 2.2, 0.95, "案 2", RED)
tb(s, x2 + 0.25, 2.6, cw2 - 0.5, 0.4, "フェーズ背骨型（推し）", size=15.5, bold=True, font=MINCHO)
tb(s, x2 + 0.25, 3.08, cw2 - 0.5, 1.0,
   "序章 → THE MODELアンサー（圧縮）→ 横の視点と「フェーズ」の発見 → フェーズ別に原則を編む（各フェーズにGTM・システム・人を織り込む）→ 終章 自分のPRINCIPLE",
   size=11, color=SOFT, spacing=16.5)
tb(s, x2 + 0.25, 4.12, cw2 - 0.5, 1.0, [
    [("○ 読者が現在地から引ける。「地図→各エリアの原則→自分の原則」と体験が一直線", {"color": BLUE})],
    [("× 横断テーマをフェーズに割り付ける再構成コストが大きい", {"color": RED})],
], size=11, spacing=16)
tb(s, MX, 5.45, CW, 1.4, [[
    ("折衷（実務案）：二部構成。", {"bold": True, "color": BLUE}),
    ("第1部「地図」＝アンサー＋横の視点＋フェーズの発見。第2部「原則」＝GTM／マネジメントシステム／人・組織のテーマ章に、全原則へ「効くフェーズ」のタグを付ける。テーマ章の書きやすさとフェーズの引きやすさを両立、現ドラフトからの移行も現実的。当日はここを一番議論したい。", {}),
]], size=12, spacing=18.5, fill=BLUESOFT, border=(BLUE, 1))

# ============ 10. Issue 4 product design ============
s = header(base(), "06 — ISSUE 4", "論点4：見せ方 —「原則」をプロダクトとして設計する")
prods = [
    ("原則の型化", "散在する「P」を〈言い切り1行＋なぜ＋実話1本＋読者への問い〉の定型に統一。全50〜70本を通し番号で管理し、巻末に原則インデックス（フェーズ×テーマ）。→ 持ち帰れる道具 その1"),
    ("ワークの本線化", "仮案の章末ワークを正式採用し、終章「自分のPRINCIPLEを作る」で回収する一冊通しの体験に。書き込んだ瞬間、本が「自分の原理原則ノート」になる。→ 道具 その2"),
    ("図版", "成長4パターン（A〜D）・フェーズ地図（STARS改）・時間差の構造は本書の顔になる図。前作の「あのスライド」級に磨く。2004年スライド等の一次資料は写真的に見せる"),
    ("引用の整理", "ワトキンス・ダリオ・ワイズマン・Playing to Win ほか引用多数。許諾要否の洗い出しを編集側で早期に実施（著者注にも【要確認】複数）。オリジナル原則との比率管理を兼ねる"),
    ("AI章の鮮度", "第1章後半のAI論は出版時点で必ず古びる。校了直前アップデート枠として設計し、構造（変わらない原則）と時事（変わる実行手段）を段落レベルで分離"),
]
y = 2.0
for it in prods:
    box = rect(s, MX, y, 2.5, 0.88, fill=SURFACE)
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = it[0]
    set_run_font(r, GOTHIC, 12.5, True, INK)
    tb(s, MX + 2.7, y, CW - 2.7, 0.88, it[1], size=11, spacing=15.5, anchor=MSO_ANCHOR.MIDDLE)
    y += 1.0

# ============ 11. Title candidates ============
s = header(base(), "07 — TITLE", "タイトルの方向感",
           "今日決めず、「THEシリーズ英語一語で行くか」の方向感だけ合意できれば十分")
rows = [
    ["候補", "根拠・狙い", "論点"],
    ["THE PRINCIPLE", "序章の結び「自分にとってのPRINCIPLEを作ってほしい」と完全整合。THEシリーズの連続性。単数形＝あなた自身の一つの原理原則という渡し方", "ダリオ『PRINCIPLES』との距離感（本文で参照しており確信犯として扱えるか）。カタカナ読みの座り"],
    ["THE PRINCIPLES", "複数の原則集であることに忠実", "ダリオとほぼ同名。検索・書店で埋没"],
    ["THE MODEL 2", "続編認知が最速。現ドラフトの作業名", "「型の続き」を期待させ、本書の主張（型ではなく原則）と自己矛盾"],
    ["THE PHASE", "フェーズ背骨案に忠実。「現在地の地図」が一語で立つ", "原理原則というテーマ性が消える。語の強度が弱い"],
]
sp = {(1, 0): (REDSOFT, RED, True), (2, 0): (SURFACE, INK, True), (3, 0): (SURFACE, INK, True), (4, 0): (SURFACE, INK, True)}
styled_table(s, MX, 2.1, rows, [2.3, 5.2, 4.43], 0.76, special_fills=sp)
tb(s, MX, 6.15, CW, 0.6,
   "推し：『THE PRINCIPLE』＋日本語サブタイトルで実利を補う（例：「なぜあの会社は、同じ壁にぶつかるのか」「フェーズで読み解く経営の原理原則」）",
   size=12, bold=True, color=BLUE)

# ============ 12. Risks ============
s = header(base(), "08 — RISKS", "正直な懸念（先に自分たちで挙げておく）")
rows = [
    ["懸念", "中身", "打ち手"],
    ["続編の宿命", "前作比で必ず評価される。「型」ほどのキャッチーさが原理原則にはない", "持ち帰れる道具（原則インデックス・ワーク）を型の代替に。帯は「答え」ではなく「地図」を約束"],
    ["外資IT特化の狭さ", "n=19の具体は強いが読者母数は狭い。「業界回顧録」に見えたら負け", "無理に汎用化せず、結論だけ汎用化（骨子の方針通り）。章末ワークが読み替え装置として機能するか編集で検証"],
    ["引用密度", "他者フレームワーク中心に見えると前作の反省と矛盾。許諾実務も重い", "主役を福田さんの「P」に。引用は出典明記の脇役へ。許諾リストを編集側で即作成"],
    ["AI章の陳腐化", "出版まで1年前後と想定するとAI記述は確実に古びる", "校了直前アップデート枠＋「変わらないもの」中心の記述に寄せる"],
    ["「関心ある人いるのかな」問題", "前作でも途中経過で陥った不安（福田さんメール）", "執筆と並行してAIdiver等で一部を連載・講演化し、読者反応を先に取る。言葉の独り歩き対策も編集が設計"],
]
styled_table(s, MX, 2.0, rows, [2.6, 4.6, 4.73], 0.84, hd_col=0)

# ============ 13. Next ============
s = header(base(), "09 — NEXT", "進め方（案）")
steps = [
    ("本日", "軸・読者・構成方針・タイトル方向感の合意。宿題の切り分け"),
    ("〜9月中旬", "編集部：合意した軸で章構成表（原則の割付マトリクス付き）を作成 → 福田さんレビュー"),
    ("9月下旬〜", "福田さん：確定構成で第1稿の執筆開始（完成度の高い終章・フェーズ章から着手する案）。編集部：引用許諾リスト・図版ラフ・ワーク設計"),
    ("執筆中", "章単位のレビューサイクル（章ごとに壁打ち。全部書いてから、にしない）。連載・講演での先行検証を並走"),
    ("校了前", "AI関連記述の最終アップデート。タイトル・帯の最終決定"),
]
y = 2.05
for st in steps:
    tb(s, MX, y, 1.9, 0.76, st[0], size=13, bold=True, color=BLUE, font="Arial", anchor=MSO_ANCHOR.MIDDLE)
    ln = s.shapes.add_connector(1, Inches(MX + 2.0), Inches(y + 0.05), Inches(MX + 2.0), Inches(y + 0.71))
    ln.line.color.rgb = rgb(LINE); ln.line.width = Pt(1)
    tb(s, MX + 2.25, y, CW - 2.25, 0.76, st[1], size=12.5, spacing=17, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.9
tb(s, MX, y + 0.15, CW, 0.4, "編集体制：押久保剛（統括）＋大久保（編集担当：丸井さん・川上さん書籍の担当編集）",
   size=11.5, color=SOFT)

out = r"C:\Users\020168\Documents\Claude\oshikubo_office\workplace\projects\the-model2-principle\THE_PRINCIPLE_方向性たたき_v1.pptx"
prs.save(out)
print("saved:", out)
